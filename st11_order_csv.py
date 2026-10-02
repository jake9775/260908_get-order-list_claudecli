# -*- coding: utf-8 -*-
"""
11번가 월별 결제 내역 엑셀 자동화

내가 지정한 연/월에 대해, Gmail로 온 11번가 결제 메일을 찾아서 엑셀(.xlsx) 파일로 정리한다.
발신자에 "11st"가 들어가고 제목에 "결제" 또는 "취소"가 들어간 메일을 검색한 뒤,
제목에 "정상 결제되었습니다"가 있으면 결제, "취소"가 있으면 취소로 처리한다.
(인증번호·고객정보 변경 같은 안내 메일은 조용히 제외한다.)

  - 메일 1통 = 엑셀 1줄 (주문번호·상품이 여러 개면 " / "로 이어서 적는다)
  - 결제일시는 메일 수신 시각(한국시간). 본문에는 주문일(날짜)만 있고 시각이 없기 때문이다.
  - 결제금액은 "총 결제금액" (할인이 반영된 실제 결제액)
  - 결제방법은 결제수단 줄 그대로 (예: 카카오페이 신용카드)
  - 상품명은 옵션 줄을 뺀 상품명만 적는다

※ 취소 메일은 실제 샘플이 없어서 형식을 확인하지 못했다. 읽지 못하면 [건너뜀]으로 표시한다.

이 프로그램은 내 Gmail 계정에서만(읽기 전용) 메일을 읽어오고, 확인한 내용은
모두 이 컴퓨터 안에서만 처리되며, 결과는 output 폴더의 엑셀 파일로만
저장된다. 외부로 전송되거나 별도로 저장되지 않는다.

사용법:
    run_11st.bat 을 더블클릭하거나, 터미널에서 `python st11_order_csv.py` 실행
"""

import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

import mapping_rules
from gmarket_order_csv import (  # 지마켓·옥션과 같은 로그인·검색·본문 변환 도구를 그대로 쓴다
    ask_year_month,
    build_query,
    get_gmail_service,
    html_to_text,
    list_message_ids,
    month_range,
)
from auction_order_csv import get_message  # 제목·본문·수신 시각(한국시간)을 돌려준다

# --- 경로 설정 -----------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
RULES_PATH = os.path.join(BASE_DIR, "rules", "memo_mapping_11번가.csv")

# 상품명 규칙 파일에서 조건으로 쓸 수 있는 칸 (지마켓·옥션과 같음)
RULE_CONDITION_COLUMNS = ("PG사", "상점명", "결제금액", "구분", "주문번호", "결제방법")
RULE_NUMBER_COLUMNS = ("결제금액",)

OUTPUT_COLUMNS = ["결제일시", "PG사", "상점명", "상품명", "결제금액", "구분", "주문번호", "결제방법"]
RIGHT_ALIGN_COLUMNS = {"결제금액"}

ST11_SENDER = "11st"
SUBJECT_FILTER = "(subject:결제 OR subject:취소)"
PG_NAME = "11번가"
STORE_NAME = "11번가"

AMOUNT_RE = r"([\d,]+)\s*원"


# --- 1. 메일 본문에서 항목 추출 -------------------------------------------
def extract_orders(lines):
    """'2026-08-07 주문번호20260807090933145' 줄마다 (주문번호, 상품명 목록)을 만든다.
    상품명은 주문번호 줄 다음 줄이고, '○○원 / ○개' 줄을 만나면 한 상품이 끝난다.
    그 사이의 나머지 줄은 옵션이라 쓰지 않는다."""
    orders = []
    i = 0
    while i < len(lines):
        m = re.search(r"주문번호\s*(\d{8,})", lines[i])
        if not m:
            i += 1
            continue
        order_no = m.group(1)
        names = []
        candidate = None  # 상품 칸의 첫 줄. 뒤에 '○○원 / ○개' 줄이 이어질 때만 상품명으로 인정한다.
        first_line_of_next = True
        i += 1
        while i < len(lines) and not re.search(r"주문번호\s*\d{8,}", lines[i]):
            if re.fullmatch(rf"{AMOUNT_RE}\s*/\s*\d+개", lines[i]):
                if candidate:
                    names.append(candidate)
                candidate = None
                first_line_of_next = True
            elif lines[i].startswith(("총 배송비", "판매자", "배송 정보")):
                break
            elif candidate is None and first_line_of_next:
                candidate = lines[i]
                first_line_of_next = False
            i += 1
        orders.append((order_no, names))
    return orders


def extract_payment(lines):
    """(총 결제금액, 결제방법)을 반환한다. 못 찾으면 None.
    '총 결제금액' 다음 줄이 금액이고, 그 뒤 '주문 상품' 전까지의 '결제수단 ○○원' 줄이 결제방법이다."""
    if "총 결제금액" not in lines:
        return None
    idx = lines.index("총 결제금액")
    if idx + 1 >= len(lines):
        return None
    m = re.fullmatch(AMOUNT_RE, lines[idx + 1])
    if not m:
        return None
    amount = int(m.group(1).replace(",", ""))

    methods = []
    for line in lines[idx + 2:]:
        if line.startswith("주문 상품"):
            break
        pm = re.fullmatch(rf"(.+?)\s*{AMOUNT_RE}", line)
        if pm:
            methods.append(pm.group(1).strip())
    return amount, " + ".join(methods)


def parse_st11_message(subject, html, received, msg_id, target_year, target_month, warn):
    # 결제도 취소도 아닌 안내 메일은 조용히 제외한다.
    if "정상 결제되었습니다" in subject:
        kind = "결제"
    elif "취소" in subject:
        kind = "취소"
    else:
        return None
    if not (received.year == target_year and received.month == target_month):
        return None

    lines = html_to_text(html).split("\n")

    orders = extract_orders(lines)
    if not orders:
        warn(f"메일(id={msg_id})에서 주문번호를 찾지 못해 건너뜁니다.")
        return None
    order_numbers = [o[0] for o in orders]

    payment = extract_payment(lines)
    if payment is None:
        warn(f"메일(id={msg_id}, 주문번호 {order_numbers[0]}, {kind})에서 결제금액을 찾지 못해 건너뜁니다.")
        return None
    amount, method = payment

    names = [name for _, product_names in orders for name in product_names]
    return {
        "_dt": received,
        "결제일시": received.strftime("%Y-%m-%d %H:%M"),
        "PG사": PG_NAME,
        "상점명": STORE_NAME,
        "상품명": " / ".join(names),
        "결제금액": amount,
        "구분": kind,
        "주문번호": " / ".join(order_numbers),
        "결제방법": method,
    }


# --- 2. 엑셀(xlsx) 저장 ------------------------------------------------------
def save_xlsx(rows, year, month):
    """정렬(왼쪽/오른쪽), 날짜·금액 서식, 헤더 필터가 적용된 엑셀 파일로 저장한다.
    주문번호는 숫자로 저장하면 지수 표기 등으로 깨질 수 있어 텍스트로 고정한다."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"{year % 100:02d}{month:02d}00_11번가결제내역_claudecli.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "결제내역"

    ws.append(OUTPUT_COLUMNS)
    for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx).alignment = Alignment(
            horizontal="right" if col_name in RIGHT_ALIGN_COLUMNS else "left"
        )

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
            value = row["_dt"] if col_name == "결제일시" else row[col_name]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(
                horizontal="right" if col_name in RIGHT_ALIGN_COLUMNS else "left"
            )
            if col_name == "결제일시":
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif col_name == "결제금액":
                cell.number_format = '#,##0"원"'
            elif col_name == "주문번호":
                cell.number_format = "@"  # 텍스트 고정 (지수 표기 방지)

    last_row = len(rows) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{last_row}"

    wb.save(filepath)
    return filepath


# --- 메인 ------------------------------------------------------------------
def main():
    print("=== 11번가 월별 결제 내역 엑셀 만들기 ===\n")

    year, month = ask_year_month()
    print(f"\n{year}년 {month}월 결제 내역을 조회합니다. 잠시만 기다려주세요...\n")

    service = get_gmail_service()
    start, end = month_range(year, month)

    warnings = []

    def warn(message):
        warnings.append(message)
        print(f"  [건너뜀] {message}")

    rows = []
    ids = list_message_ids(service, build_query(start, end, ST11_SENDER, SUBJECT_FILTER))
    print(f"11번가 대상 메일 {len(ids)}건")
    for msg_id in ids:
        subject, html, received = get_message(service, msg_id)
        if not html:
            warn(f"메일(id={msg_id})의 내용을 읽을 수 없어 건너뜁니다.")
            continue
        row = parse_st11_message(subject, html, received, msg_id, year, month, warn)
        if row:
            rows.append(row)

    rows.sort(key=lambda r: r["_dt"], reverse=True)  # 결제일시 내림차순

    # 상품명이 빈칸인 행에 규칙 파일의 상품명을 채운다.
    mapping_rules.ensure_template(RULES_PATH, RULE_CONDITION_COLUMNS + ("상품명",))
    rules = mapping_rules.load(RULES_PATH, RULE_CONDITION_COLUMNS, RULE_NUMBER_COLUMNS)
    filled = mapping_rules.apply(rows, rules, RULE_NUMBER_COLUMNS)
    if filled:
        print(f"상품명 규칙으로 {filled}건을 채웠습니다.")

    filepath = save_xlsx(rows, year, month)

    print()
    if rows:
        print(f"총 {len(rows)}건을 정리해서 저장했습니다.")
    else:
        print(f"{year}년 {month}월에는 11번가 결제 내역이 없습니다. (항목만 있는 빈 엑셀 파일을 만들었습니다)")
    print(f"저장 위치: {filepath}")

    if warnings:
        print(f"\n※ {len(warnings)}건은 형식을 해석하지 못해 건너뛰었습니다. 위 [건너뜀] 내용을 확인해주세요.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n취소되었습니다.")
    except PermissionError as e:
        print(f"\n[오류] 파일에 저장하지 못했습니다: {e.filename}")
        print("→ 이 엑셀 파일을 이미 열어놓은 상태라서 저장이 막혔을 가능성이 매우 높습니다.")
        print("   해당 파일을 닫은 뒤, 프로그램을 다시 실행해주세요.")
    except Exception as e:  # noqa: BLE001 - 비개발자용 프로그램이므로 원인을 그대로 보여준다
        print(f"\n[오류] 문제가 발생했습니다: {e}")
    finally:
        input("\n엔터 키를 누르면 창이 닫힙니다...")
