# -*- coding: utf-8 -*-
"""
옥션 월별 결제 내역 엑셀 자동화

내가 지정한 연/월에 대해, Gmail로 온 옥션 주문 메일을 찾아서 엑셀(.xlsx) 파일로 정리한다.
발신자는 auction@auction.co.kr 이고, 제목에 "주문"과 "감사합니다"가 들어간 메일
("주문해 주셔서 감사합니다")만 처리한다.

  - 메일 1통 = 엑셀 1줄 (주문번호가 여러 개면 " / "로 이어서 적는다)
  - 결제일시는 메일 수신 시각(한국시간). 본문에는 날짜만 있고 시각이 없기 때문이다.
  - 결제금액은 "결제 정보"의 Smile Pay + 스마일캐시(·머니) 합계 (실제 결제액)
  - 상품명은 주문 상품 전부를 " / "로 이어서 적는다 (옵션은 적지 않음)

취소/환불 메일은 처리하지 않는다. 같은 주문번호가 이미 처리됐으면 중복으로 보고 제외한다.

이 프로그램은 내 Gmail 계정에서만(읽기 전용) 메일을 읽어오고, 확인한 내용은
모두 이 컴퓨터 안에서만 처리되며, 결과는 output 폴더의 엑셀 파일로만
저장된다. 외부로 전송되거나 별도로 저장되지 않는다.

사용법:
    run_auction.bat 을 더블클릭하거나, 터미널에서 `python auction_order_csv.py` 실행
"""

import os
import re
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

import mapping_rules
from gmarket_order_csv import (  # 지마켓과 같은 로그인·검색·본문 해석 도구를 그대로 쓴다
    ask_year_month,
    build_query,
    get_gmail_service,
    html_to_text,
    list_message_ids,
    month_range,
    _find_part,
)

# --- 경로 설정 -----------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
RULES_PATH = os.path.join(BASE_DIR, "rules", "memo_mapping_옥션.csv")

# 상품명 규칙 파일에서 조건으로 쓸 수 있는 칸 (지마켓과 같음)
RULE_CONDITION_COLUMNS = ("PG사", "상점명", "결제금액", "구분", "주문번호", "결제방법")
RULE_NUMBER_COLUMNS = ("결제금액",)

OUTPUT_COLUMNS = ["결제일시", "PG사", "상점명", "상품명", "결제금액", "구분", "주문번호", "결제방법"]
RIGHT_ALIGN_COLUMNS = {"결제금액"}

AUCTION_SENDER = "auction@auction.co.kr"
SUBJECT_FILTER = "subject:주문 subject:감사합니다"
PG_NAME = "스마일페이"
STORE_NAME = "옥션"
KST = timezone(timedelta(hours=9))


# --- 1. Gmail 메일 읽기 ----------------------------------------------------
def get_message(service, msg_id):
    """제목, HTML 본문, 수신 시각(한국시간)을 반환한다."""
    msg = service.users().messages().get(userId="me", id=msg_id, format="full").execute()
    payload = msg.get("payload", {})
    headers = {h["name"]: h["value"] for h in payload.get("headers", [])}
    subject = headers.get("Subject", "")
    body = _find_part(payload, "text/html")
    if body is None:
        body = _find_part(payload, "text/plain")
    received = datetime.fromtimestamp(int(msg["internalDate"]) / 1000, tz=KST).replace(tzinfo=None)
    return subject, body or "", received


# --- 2. 메일 본문에서 항목 추출 -------------------------------------------
def extract_order_numbers(text):
    """본문의 모든 '주문번호 숫자'를 순서대로 꺼낸다 (중복 제거)."""
    found = re.findall(r"주문번호\s+(\d{6,})", text)
    return list(dict.fromkeys(found))


def extract_product_names(html):
    """상품 링크(주소에 t=vp 포함)의 글자가 상품명이다. 옵션은 링크가 아니라서 제외된다."""
    names = []
    for m in re.finditer(r"""<a\b[^>]*href=["'][^"']*t=vp[^"']*["'][^>]*>(.*?)</a>""", html, re.DOTALL):
        name = html_to_text(m.group(1)).replace("\n", " ").strip()
        if name and name not in names:
            names.append(name)
    return names


def extract_payment(text):
    """'결제 정보' 칸에서 '상품금액 합계' 앞까지의 (이름, 금액) 줄을 읽어
    (결제금액 합계, 결제방법)을 반환한다. 못 찾으면 None."""
    lines = text.split("\n")
    if "결제 정보" not in lines:
        return None
    # 예전 메일은 "Smile / Pay"처럼 이름이 두 줄로 나뉘어 있어서, 구간 전체를 한 줄로 이어서 읽는다.
    section = []
    for line in lines[lines.index("결제 정보") + 1:]:
        if line.startswith("상품금액 합계"):
            break
        section.append(line)
    pairs = re.findall(r"(\D+?)\s*([\d,]+)\s*원", " ".join(section))
    if not pairs:
        return None
    total = 0
    methods = []
    for label, value in pairs:
        amount = int(value.replace(",", ""))
        total += amount
        if amount > 0:
            methods.append(" ".join(label.split()))

    # 무통장입금처럼 같은 금액이 두 줄(입금 + 스마일캐시·머니)로 나오는 메일은 위 합계가
    # 2배가 되므로, 하단의 "상품금액 합계 + 배송비 합계 - 총 할인금액"을 실결제액으로 우선한다.
    tail = text[text.index("결제 정보"):]
    parts = [
        re.search(label + r"[^\n]*\n([\d,]+) 원", tail)
        for label in ("상품금액 합계", "배송비 합계", "총 할인금액")
    ]
    if all(parts):
        goods, shipping, discount = (int(m.group(1).replace(",", "")) for m in parts)
        total = goods + shipping - discount
    return total, " + ".join(methods)


def parse_auction_message(subject, html, received, msg_id, target_year, target_month, seen_orders, warn):
    # 주문 안내가 아닌 메일은 조용히 제외한다.
    if "주문" not in subject or "감사합니다" not in subject:
        return None
    if not (received.year == target_year and received.month == target_month):
        return None

    text = html_to_text(html)

    orders = extract_order_numbers(text)
    if not orders:
        warn(f"메일(id={msg_id})에서 주문번호를 찾지 못해 건너뜁니다.")
        return None
    if all(o in seen_orders for o in orders):
        return None  # 이미 처리한 주문번호
    payment = extract_payment(text)
    if payment is None:
        warn(f"메일(id={msg_id}, 주문번호 {orders[0]})에서 결제금액을 찾지 못해 건너뜁니다.")
        return None
    amount, method = payment
    seen_orders.update(orders)

    return {
        "_dt": received,
        "결제일시": received.strftime("%Y-%m-%d %H:%M"),
        "PG사": PG_NAME,
        "상점명": STORE_NAME,
        "상품명": " / ".join(extract_product_names(html)),
        "결제금액": amount,
        "구분": "결제",
        "주문번호": " / ".join(orders),
        "결제방법": method,
    }


# --- 3. 엑셀(xlsx) 저장 ------------------------------------------------------
def save_xlsx(rows, year, month):
    """정렬(왼쪽/오른쪽), 날짜·금액 서식, 헤더 필터가 적용된 엑셀 파일로 저장한다.
    주문번호는 숫자로 저장하면 지수 표기 등으로 깨질 수 있어 텍스트로 고정한다."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"{year % 100:02d}{month:02d}00_옥션결제내역_claudecli.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "결제내역"

    ws.append(OUTPUT_COLUMNS)
    for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx).alignment = Alignment(
            horizontal="right" if col_name in RIGHT_ALIGN_COLUMNS else "left"
        )

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
            value = row["_dt"] if col_name == "결제일시" else row[col_name]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(
                horizontal="right" if col_name in RIGHT_ALIGN_COLUMNS else "left"
            )
            if col_name == "결제일시":
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif col_name == "결제금액":
                cell.number_format = '#,##0"원"'
            elif col_name == "주문번호":
                cell.number_format = "@"  # 텍스트 고정 (지수 표기 방지)

    last_row = len(rows) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{last_row}"

    wb.save(filepath)
    return filepath


# --- 메인 ------------------------------------------------------------------
def main():
    print("=== 옥션 월별 결제 내역 엑셀 만들기 ===\n")

    year, month = ask_year_month()
    print(f"\n{year}년 {month}월 결제 내역을 조회합니다. 잠시만 기다려주세요...\n")

    service = get_gmail_service()
    start, end = month_range(year, month)

    warnings = []

    def warn(message):
        warnings.append(message)
        print(f"  [건너뜀] {message}")

    rows = []
    seen_orders = set()
    ids = list_message_ids(service, build_query(start, end, AUCTION_SENDER, SUBJECT_FILTER))
    print(f"옥션 대상 메일 {len(ids)}건")
    for msg_id in ids:
        subject, html, received = get_message(service, msg_id)
        if not html:
            warn(f"메일(id={msg_id})의 내용을 읽을 수 없어 건너뜁니다.")
            continue
        row = parse_auction_message(subject, html, received, msg_id, year, month, seen_orders, warn)
        if row:
            rows.append(row)

    rows.sort(key=lambda r: r["_dt"], reverse=True)  # 결제일시 내림차순

    # 상품명이 빈칸인 행에 규칙 파일의 상품명을 채운다.
    mapping_rules.ensure_template(RULES_PATH, RULE_CONDITION_COLUMNS + ("상품명",))
    rules = mapping_rules.load(RULES_PATH, RULE_CONDITION_COLUMNS, RULE_NUMBER_COLUMNS)
    filled = mapping_rules.apply(rows, rules, RULE_NUMBER_COLUMNS)
    if filled:
        print(f"상품명 규칙으로 {filled}건을 채웠습니다.")

    filepath = save_xlsx(rows, year, month)

    print()
    if rows:
        print(f"총 {len(rows)}건을 정리해서 저장했습니다.")
    else:
        print(f"{year}년 {month}월에는 옥션 결제 내역이 없습니다. (항목만 있는 빈 엑셀 파일을 만들었습니다)")
    print(f"저장 위치: {filepath}")

    if warnings:
        print(f"\n※ {len(warnings)}건은 형식을 해석하지 못해 건너뛰었습니다. 위 [건너뜀] 내용을 확인해주세요.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n취소되었습니다.")
    except PermissionError as e:
        print(f"\n[오류] 파일에 저장하지 못했습니다: {e.filename}")
        print("→ 이 엑셀 파일을 이미 열어놓은 상태라서 저장이 막혔을 가능성이 매우 높습니다.")
        print("   해당 파일을 닫은 뒤, 프로그램을 다시 실행해주세요.")
    except Exception as e:  # noqa: BLE001 - 비개발자용 프로그램이므로 원인을 그대로 보여준다
        print(f"\n[오류] 문제가 발생했습니다: {e}")
    finally:
        input("\n엔터 키를 누르면 창이 닫힙니다...")
