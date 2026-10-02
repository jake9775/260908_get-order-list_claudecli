# -*- coding: utf-8 -*-
"""
지마켓 월별 결제 내역 엑셀 자동화

내가 지정한 연/월에 대해, Gmail로 온 G마켓 결제완료 메일을 찾아서
엑셀(.xlsx) 파일로 정리한다. 발신자는 mailmaster@corp.gmarket.co.kr 하나뿐이고,
제목에 "결제"가 들어간 메일을 검색한 뒤 제목에 "결제가 완료"가 있는 메일만
처리한다. (같은 검색에 걸리는 "스마일페이 결제사업부문 분할 안내" 같은 공지
메일은 조용히 제외한다.)

  - 주문(장바구니) 1건 = 엑셀 1줄
  - 결제금액은 "총 결제금액" (스마일캐시·머니 사용분을 뺀 실제 결제액)
  - 상품명은 주문 상품 전부를 " / "로 이어서 적는다 (옵션은 적지 않음)

2020년 4월 17일 이후의 메일 형식만 처리한다. 그 이전 메일은 형식이 완전히
달라서 [건너뜀]으로 표시된다. 취소/환불 메일은 처리하지 않는다.

이 프로그램은 내 Gmail 계정에서만(읽기 전용) 메일을 읽어오고, 확인한 내용은
모두 이 컴퓨터 안에서만 처리되며, 결과는 output 폴더의 엑셀 파일로만
저장된다. 외부로 전송되거나 별도로 저장되지 않는다.

사용법:
    run_gmarket.bat 을 더블클릭하거나, 터미널에서 `python gmarket_order_csv.py` 실행
"""

import base64
import os
import re
import sys
from datetime import datetime, timedelta
from html.parser import HTMLParser

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

import mapping_rules

# --- 경로 설정 -----------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CREDENTIALS_PATH = os.path.join(BASE_DIR, "credentials.json")
TOKEN_PATH = os.path.join(BASE_DIR, "token.json")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
RULES_PATH = os.path.join(BASE_DIR, "rules", "memo_mapping_지마켓.csv")

# 상품명 규칙 파일에서 조건으로 쓸 수 있는 칸 (작성법은 mapping_rules.py 참고)
RULE_CONDITION_COLUMNS = ("PG사", "상점명", "결제금액", "구분", "주문번호", "결제방법")
RULE_NUMBER_COLUMNS = ("결제금액",)

# 지메일 읽기 전용 권한만 요청 (메일 삭제/발송 등은 하지 않음)
SCOPES = ["https://www.googleapis.com/auth/gmail.readonly"]

# 엑셀 파일에 실제로 저장할 컬럼 순서 (구글플레이와 같음)
OUTPUT_COLUMNS = ["결제일시", "PG사", "상점명", "상품명", "결제금액", "구분", "주문번호", "결제방법"]
RIGHT_ALIGN_COLUMNS = {"결제금액"}

GMARKET_SENDER = "mailmaster@corp.gmarket.co.kr"
SUBJECT_FILTER = "subject:결제"
PG_NAME = "스마일페이"
STORE_NAME = "G마켓"


# --- 1. 구글 로그인 / 인증 ------------------------------------------------
def get_gmail_service():
    """최초 1회만 브라우저 로그인이 뜨고, 이후에는 저장된 로그인 정보를 재사용한다."""
    if not os.path.exists(CREDENTIALS_PATH):
        print(f"[오류] '{os.path.basename(CREDENTIALS_PATH)}' 파일을 찾을 수 없습니다.")
        print("SETUP.md 안내를 참고해서 구글 클라우드 콘솔에서 인증 파일을 받아")
        print(f"이 폴더({BASE_DIR})에 넣어주세요.")
        sys.exit(1)

    creds = None
    if os.path.exists(TOKEN_PATH):
        creds = Credentials.from_authorized_user_file(TOKEN_PATH, SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_PATH, SCOPES)
            creds = flow.run_local_server(port=0)
        with open(TOKEN_PATH, "w", encoding="utf-8") as f:
            f.write(creds.to_json())

    return build("gmail", "v1", credentials=creds)


# --- 2. 사용자에게 연/월 물어보기 -----------------------------------------
def ask_year_month():
    now = datetime.now()
    while True:
        raw_year = input(f"조회할 연도를 입력하세요 (예: {now.year}, 그냥 엔터=올해): ").strip()
        year_str = str(now.year) if raw_year == "" else raw_year

        raw_month = input("조회할 월을 입력하세요 (1~12): ").strip()

        if not (year_str.isdigit() and raw_month.isdigit()):
            print("→ 숫자로 입력해주세요. 다시 입력해주세요.\n")
            continue

        year = int(year_str)
        month = int(raw_month)

        if not (2000 <= year <= 2100) or not (1 <= month <= 12):
            print("→ 연도 또는 월 범위가 올바르지 않습니다. 다시 입력해주세요.\n")
            continue

        return year, month


# --- 3. Gmail 검색 --------------------------------------------------------
def month_range(year, month):
    start = datetime(year, month, 1)
    end = datetime(year + 1, 1, 1) if month == 12 else datetime(year, month + 1, 1)
    return start, end


def build_query(start, end, sender, subject_filter):
    # 시간대 경계에서 메일이 하루 밀리는 경우를 대비해 앞뒤로 하루씩 여유를 두고 검색한다.
    # 실제 대상 월인지는 본문의 결제완료 시각으로 다시 확인한다.
    q_start = (start - timedelta(days=1)).strftime("%Y/%m/%d")
    q_end = (end + timedelta(days=1)).strftime("%Y/%m/%d")
    return f"from:{sender} after:{q_start} before:{q_end} {subject_filter}"


def list_message_ids(service, query):
    ids = []
    page_token = None
    while True:
        resp = (
            service.users()
            .messages()
            .list(userId="me", q=query, pageToken=page_token, maxResults=100)
            .execute()
        )
        ids.extend(m["id"] for m in resp.get("messages", []))
        page_token = resp.get("nextPageToken")
        if not page_token:
            break
    return ids


def _decode_body_data(data):
    raw = base64.urlsafe_b64decode(data.encode("ASCII") + b"===")
    for encoding in ("utf-8", "euc-kr", "cp949"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _find_part(payload, mime_type):
    if payload.get("mimeType") == mime_type and "data" in payload.get("body", {}):
        return _decode_body_data(payload["body"]["data"])
    for part in payload.get("parts") or []:
        found = _find_part(part, mime_type)
        if found is not None:
            return found
    return None


def get_message(service, msg_id):
    """G마켓 메일은 HTML 본문만 있어서 HTML을 그대로 반환한다."""
    msg = service.users().messages().get(userId="me", id=msg_id, format="full").execute()
    payload = msg.get("payload", {})
    headers = {h["name"]: h["value"] for h in payload.get("headers", [])}
    subject = headers.get("Subject", "")
    body = _find_part(payload, "text/html")
    if body is None:
        body = _find_part(payload, "text/plain")
    return subject, body or ""


# --- 4. 메일 본문에서 항목 추출 -------------------------------------------
class _TextExtractor(HTMLParser):
    """HTML 표의 칸/줄 경계마다 줄바꿈을 넣어서, 화면에 보이는 글자를
    한 줄에 하나씩 꺼낸다."""

    BLOCK_TAGS = {"br", "p", "div", "tr", "td", "th", "li", "table", "tbody"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("style", "script", "head"):
            self.skip += 1
        if tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("style", "script", "head"):
            self.skip = max(0, self.skip - 1)
        if tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def html_to_text(html):
    parser = _TextExtractor()
    parser.feed(html)
    parser.close()
    text = "".join(parser.parts).replace("\xa0", " ")
    lines = (re.sub(r"\s+", " ", line).strip() for line in text.split("\n"))
    return "\n".join(line for line in lines if line)


def parse_paid_datetime(text):
    # 예: 2026년 9월 20일(일) 12시 12분 결제완료
    m = re.search(
        r"(\d{4})년\s*(\d{1,2})월\s*(\d{1,2})일\s*\([^)]*\)\s*(\d{1,2})시\s*(\d{1,2})분\s*결제완료",
        text,
    )
    if not m:
        return None
    try:
        return datetime(*(int(g) for g in m.groups()))
    except ValueError:
        return None


def extract_product_names(html):
    """상품마다 썸네일 칸 바로 뒤의 첫 번째 링크가 상품명이다.
    (그 다음 링크는 색상·사이즈 같은 옵션이라 쓰지 않는다)"""
    names = []
    for segment in html.split("<!-- 썸네일 -->")[1:]:
        m = re.search(r"<a\b[^>]*>(.*?)</a>", segment, re.DOTALL)
        if m:
            name = html_to_text(m.group(1)).replace("\n", " ")
            if name:
                names.append(name)
    return names


def parse_gmarket_message(subject, html, msg_id, target_year, target_month, warn):
    # 공지 메일 등 결제완료가 아닌 메일은 조용히 제외한다.
    if "결제가 완료" not in subject:
        return None

    text = html_to_text(html)

    dt = parse_paid_datetime(text)
    if dt is None:
        warn(f"메일(id={msg_id})에서 결제완료 시각을 해석하지 못해 건너뜁니다. (2020년 4월 이전 형식일 수 있음)")
        return None
    if not (dt.year == target_year and dt.month == target_month):
        return None

    # 본문 맨 위: "결제금액(또는 총 결제금액) / 62,270 / 원 / Smile Pay / …결제완료"
    m = re.search(r"결제금액\n([\d,]+)\n원\n(.*?)\n\d{4}년[^\n]*결제완료", text, re.DOTALL)
    if not m:
        warn(f"메일(id={msg_id})에서 결제금액을 찾지 못해 건너뜁니다.")
        return None
    amount = int(m.group(1).replace(",", ""))
    payment_method = " ".join(m.group(2).split("\n"))

    cart = re.search(r"장바구니 번호:\s*(\d+)", text)

    return {
        "_dt": dt,
        "결제일시": dt.strftime("%Y-%m-%d %H:%M"),
        "PG사": PG_NAME,
        "상점명": STORE_NAME,
        "상품명": " / ".join(extract_product_names(html)),
        "결제금액": amount,
        "구분": "결제",
        "주문번호": cart.group(1) if cart else "",
        "결제방법": payment_method,
    }


# --- 5. 엑셀(xlsx) 저장 ------------------------------------------------------
def save_xlsx(rows, year, month):
    """정렬(왼쪽/오른쪽), 날짜·금액 서식, 헤더 필터가 적용된 엑셀 파일로 저장한다.
    주문번호는 숫자로 저장하면 지수 표기 등으로 깨질 수 있어 텍스트로 고정한다."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"{year % 100:02d}{month:02d}00_지마켓결제내역_claudecli.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "결제내역"

    ws.append(OUTPUT_COLUMNS)
    for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx).alignment = Alignment(
            horizontal="right" if col_name in RIGHT_ALIGN_COLUMNS else "left"
        )

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
            value = row["_dt"] if col_name == "결제일시" else row[col_name]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(
                horizontal="right" if col_name in RIGHT_ALIGN_COLUMNS else "left"
            )
            if col_name == "결제일시":
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif col_name == "결제금액":
                cell.number_format = '#,##0"원"'
            elif col_name == "주문번호":
                cell.number_format = "@"  # 텍스트 고정 (지수 표기 방지)

    last_row = len(rows) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{last_row}"

    wb.save(filepath)
    return filepath


# --- 메인 ------------------------------------------------------------------
def main():
    print("=== 지마켓 월별 결제 내역 엑셀 만들기 ===\n")

    year, month = ask_year_month()
    print(f"\n{year}년 {month}월 결제 내역을 조회합니다. 잠시만 기다려주세요...\n")

    service = get_gmail_service()
    start, end = month_range(year, month)

    warnings = []

    def warn(message):
        warnings.append(message)
        print(f"  [건너뜀] {message}")

    rows = []
    ids = list_message_ids(service, build_query(start, end, GMARKET_SENDER, SUBJECT_FILTER))
    print(f"G마켓 대상 메일 {len(ids)}건")
    for msg_id in ids:
        subject, html = get_message(service, msg_id)
        if not html:
            warn(f"메일(id={msg_id})의 내용을 읽을 수 없어 건너뜁니다.")
            continue
        row = parse_gmarket_message(subject, html, msg_id, year, month, warn)
        if row:
            rows.append(row)

    rows.sort(key=lambda r: r["_dt"], reverse=True)  # 결제일시 내림차순

    # 상품명이 빈칸인 행에 규칙 파일의 상품명을 채운다.
    mapping_rules.ensure_template(RULES_PATH, RULE_CONDITION_COLUMNS + ("상품명",))
    rules = mapping_rules.load(RULES_PATH, RULE_CONDITION_COLUMNS, RULE_NUMBER_COLUMNS)
    filled = mapping_rules.apply(rows, rules, RULE_NUMBER_COLUMNS)
    if filled:
        print(f"상품명 규칙으로 {filled}건을 채웠습니다.")

    filepath = save_xlsx(rows, year, month)

    print()
    if rows:
        print(f"총 {len(rows)}건을 정리해서 저장했습니다.")
    else:
        print(f"{year}년 {month}월에는 지마켓 결제 내역이 없습니다. (항목만 있는 빈 엑셀 파일을 만들었습니다)")
    print(f"저장 위치: {filepath}")

    if warnings:
        print(f"\n※ {len(warnings)}건은 형식을 해석하지 못해 건너뛰었습니다. 위 [건너뜀] 내용을 확인해주세요.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n취소되었습니다.")
    except PermissionError as e:
        print(f"\n[오류] 파일에 저장하지 못했습니다: {e.filename}")
        print("→ 이 엑셀 파일을 이미 열어놓은 상태라서 저장이 막혔을 가능성이 매우 높습니다.")
        print("   해당 파일을 닫은 뒤, 프로그램을 다시 실행해주세요.")
    except Exception as e:  # noqa: BLE001 - 비개발자용 프로그램이므로 원인을 그대로 보여준다
        print(f"\n[오류] 문제가 발생했습니다: {e}")
    finally:
        input("\n엔터 키를 누르면 창이 닫힙니다...")
