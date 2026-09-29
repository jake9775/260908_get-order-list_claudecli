# -*- coding: utf-8 -*-
"""
구글플레이 월별 결제 내역 엑셀 자동화

내가 지정한 연/월에 대해, Gmail로 온 Google Play 영수증 메일을 찾아서
엑셀(.xlsx) 파일로 정리한다. 발신자는 googleplay-noreply@google.com 하나뿐이고,
제목에 "영수증"이 들어간 메일만 검색해서 아래 2가지만 처리한다.

  A. 결제 영수증  - 제목: "Google Play 주문 영수증(YYYY. M. D.)"
  B. 취소 영수증  - 제목: "Google Play 주문 취소 영수증(YYYY. M. D.)"

추가로 Google 결제 센터(payments-noreply@google.com)에서 온, 제목에 "결제"가
들어간 메일도 검색해서 아래 1가지를 처리한다.

  C. 결제 완료    - 제목: "Google Cloud Platform & APIs: 결제 완료" 등
                    본문에 시각이 없어 결제일시는 "메일 수신 시각"으로 기록한다.
                    같은 검색에 걸리는 "결제 수단 업데이트됨", "결제 계정 폐쇄
                    알림"은 금액이 없는 안내 메일이라 조용히 제외한다.

아래 메일들은 이번 버전에서 일부러 처리하지 않는다 (다음 버전 과제):
  - "Google Play 환불이 승인됨" - 문장형 별도 환불 메일. 위 A/B와 본문 구조가
    완전히 달라서(표 형태가 아님, 주문 날짜/결제 방법 항목 자체가 없음) 별도
    파서가 필요함. 이 메일만 단독으로 오고 취소 영수증이 안 오는 환불 건은
    이번 버전 엑셀에 잡히지 않는다.
  - "OOO 정기 결제가 취소됩니다" 같은 예정 취소 알림 - 아직 돈이 움직이지
    않은 안내 메일이라 금액 자체가 없음. 제목에 "영수증"이 없어 검색에도
    안 걸림.
  - "환불 요청을 검토하는 중" - 아직 확정되지 않은 건이라 제외.

이 프로그램은 내 Gmail 계정에서만(읽기 전용) 메일을 읽어오고, 확인한 내용은
모두 이 컴퓨터 안에서만 처리되며, 결과는 output 폴더의 엑셀 파일로만
저장된다. 외부로 전송되거나 별도로 저장되지 않는다.

사용법:
    python google_play_order_csv.py 실행
"""

import base64
import os
import re
import sys
from datetime import datetime, timedelta

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

# --- 경로 설정 -----------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CREDENTIALS_PATH = os.path.join(BASE_DIR, "credentials.json")
TOKEN_PATH = os.path.join(BASE_DIR, "token.json")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

# 지메일 읽기 전용 권한만 요청 (메일 삭제/발송 등은 하지 않음)
SCOPES = ["https://www.googleapis.com/auth/gmail.readonly"]

# 엑셀 파일에 실제로 저장할 컬럼 순서
OUTPUT_COLUMNS = ["결제일시", "PG사", "상점명", "상품명", "결제금액", "구분", "주문번호", "결제방법"]
# 금액만 오른쪽 정렬, 나머지 컬럼은 모두 왼쪽 정렬한다.
RIGHT_ALIGN_COLUMNS = {"결제금액"}

# Google Play 영수증 메일은 이 발신자 하나뿐이다.
GOOGLE_PLAY_SENDER = "googleplay-noreply@google.com"

# "영수증"이 제목에 들어간 메일만 검색하면 결제(A)·취소 영수증(B) 메일만
# 걸리고, "환불이 승인됨"이나 "정기 결제가 취소됩니다" 같은 다른 형식의
# 메일은 자동으로 제외된다.
SUBJECT_FILTER = "subject:영수증"

# Google 결제 센터(payments-noreply)는 Google Cloud 등 Play 외 서비스의 결제
# 완료 메일을 보낸다. 제목에 "결제"가 들어간 메일을 검색하되, 같은 조건에
# "결제 수단 업데이트됨"·"결제 계정 폐쇄 알림"처럼 금액이 없는 안내 메일도
# 걸리므로, 본문에 "…의 결제 금액이 …에 적용" 문장이 있는 메일만 담는다.
GOOGLE_PAYMENTS_SENDER = "payments-noreply@google.com"
PAYMENTS_SUBJECT_FILTER = "subject:결제"

# 영어 날짜 형식(예전 메일)을 해석할 때 쓰는 월 이름 표
MONTH_NAMES = {
    "Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6,
    "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12,
}


# --- 1. 구글 로그인 / 인증 ------------------------------------------------
def get_gmail_service():
    """최초 1회만 브라우저 로그인이 뜨고, 이후에는 저장된 로그인 정보를 재사용한다."""
    if not os.path.exists(CREDENTIALS_PATH):
        print(f"[오류] '{os.path.basename(CREDENTIALS_PATH)}' 파일을 찾을 수 없습니다.")
        print("SETUP.md 안내를 참고해서 구글 클라우드 콘솔에서 인증 파일을 받아")
        print(f"이 폴더({BASE_DIR})에 넣어주세요.")
        sys.exit(1)

    creds = None
    if os.path.exists(TOKEN_PATH):
        creds = Credentials.from_authorized_user_file(TOKEN_PATH, SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_PATH, SCOPES)
            creds = flow.run_local_server(port=0)
        with open(TOKEN_PATH, "w", encoding="utf-8") as f:
            f.write(creds.to_json())

    return build("gmail", "v1", credentials=creds)


# --- 2. 사용자에게 연/월 물어보기 -----------------------------------------
def ask_year_month():
    now = datetime.now()
    while True:
        raw_year = input(f"조회할 연도를 입력하세요 (예: {now.year}, 그냥 엔터=올해): ").strip()
        year_str = str(now.year) if raw_year == "" else raw_year

        raw_month = input("조회할 월을 입력하세요 (1~12): ").strip()

        if not (year_str.isdigit() and raw_month.isdigit()):
            print("→ 숫자로 입력해주세요. 다시 입력해주세요.\n")
            continue

        year = int(year_str)
        month = int(raw_month)

        if not (2000 <= year <= 2100) or not (1 <= month <= 12):
            print("→ 연도 또는 월 범위가 올바르지 않습니다. 다시 입력해주세요.\n")
            continue

        return year, month


# --- 3. Gmail 검색 --------------------------------------------------------
def month_range(year, month):
    start = datetime(year, month, 1)
    end = datetime(year + 1, 1, 1) if month == 12 else datetime(year, month + 1, 1)
    return start, end


def build_query(start, end, sender, subject_filter):
    # 시간대 경계에서 메일이 하루 밀리는 경우를 대비해 앞뒤로 하루씩 여유를 두고 검색한다.
    # 실제 대상 월인지는 메일 날짜를 다시 확인해 한번 더 정확히 걸러낸다.
    q_start = (start - timedelta(days=1)).strftime("%Y/%m/%d")
    q_end = (end + timedelta(days=1)).strftime("%Y/%m/%d")
    return f"from:{sender} after:{q_start} before:{q_end} {subject_filter}"


def list_message_ids(service, query):
    ids = []
    page_token = None
    while True:
        resp = (
            service.users()
            .messages()
            .list(userId="me", q=query, pageToken=page_token, maxResults=100)
            .execute()
        )
        ids.extend(m["id"] for m in resp.get("messages", []))
        page_token = resp.get("nextPageToken")
        if not page_token:
            break
    return ids


def _decode_body_data(data):
    raw = base64.urlsafe_b64decode(data.encode("ASCII") + b"===")
    for encoding in ("utf-8", "euc-kr", "cp949"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _find_part(payload, mime_type):
    if payload.get("mimeType") == mime_type and "data" in payload.get("body", {}):
        return _decode_body_data(payload["body"]["data"])
    for part in payload.get("parts") or []:
        found = _find_part(part, mime_type)
        if found is not None:
            return found
    return None


def get_message(service, msg_id):
    """Google Play 영수증 메일은 plain text 본문이 "주문 번호:" 같은 라벨
    형태로 깔끔하게 정리되어 있어서, HTML보다 text/plain을 우선 사용한다.
    구분(결제/취소) 판단에 제목도 필요해서 제목과 본문을 함께 반환한다.
    결제 센터 메일은 본문에 시각이 없어서, 메일 수신 시각(이 PC 현지시간)도
    함께 반환한다."""
    msg = service.users().messages().get(userId="me", id=msg_id, format="full").execute()
    payload = msg.get("payload", {})
    headers = {h["name"]: h["value"] for h in payload.get("headers", [])}
    subject = headers.get("Subject", "")
    body = _find_part(payload, "text/plain")
    if body is None:
        body = _find_part(payload, "text/html")
    received_dt = datetime.fromtimestamp(int(msg["internalDate"]) / 1000)
    return subject, body or "", received_dt


# --- 4. 메일 본문에서 항목 추출 -------------------------------------------
def extract_order_number(body):
    m = re.search(r"(?:주문\s*번호|Order\s*number):\s*(\S+)", body)
    return m.group(1) if m else ""


def extract_order_datetime_text(body):
    m = re.search(r"(?:주문\s*날짜|Order\s*date):\s*(.+)", body)
    return m.group(1).strip() if m else ""


def extract_store_name(body):
    """상점명(판매자/개발자명)이 문장 속에 섞여 있어서, 메일에 쓰이는
    문장 패턴 여러 개를 순서대로 시도해서 뽑아낸다."""
    patterns = [
        r"Google Play의\s*(.+?)에서\s*구매하(?:셨습니다|신 내역이 취소되었습니다)",
        r"Google Play를 통한\s*(.+?)\s*구독이\s*갱신",
        r"Google Play에서\s*(.+?)의\s*구독이\s*갱신",
        r"Google Play에서\s*(.+?)의\s*정기\s*결제를\s*시작하셨습니다",  # 신규 구독 시작
        r"Google Play에서\s*(.+?)의\s*무료\s*체험",  # 무료 체험 시작
        r"Google Play\s+(.+?)\s*구독\s*구매를\s*업데이트했습니다",  # 구독(결제수단) 업데이트
        # 예전 영어 형식 메일
        r"purchase from\s*(.+?)\s*on Google Play has been (?:canceled|cancelled)",
        r"subscription from\s*(.+?)\s*on Google Play has renewed",
    ]
    for pattern in patterns:
        m = re.search(pattern, body)
        if m:
            return m.group(1).strip()
    return ""


def extract_item_and_amount(body):
    """"상품 가격"/"상품 환불 금액"(최근, 한글) 또는 "Item Price"/"Item Amount
    refunded"(예전, 영어) 항목 아래에서 상품명과 금액을 뽑는다. 구독 상품은
    금액 앞에 "매월" 같은 글자가 붙기도 해서 무시하고 숫자만 뽑는다.

    금액 앞의 통화 기호는 보통 "₩"이지만, 실제로 받은 메일 중 일부는 원문
    자체에 "₩" 대신 "a" 같은 엉뚱한 글자가 들어있는 경우가 있었다(Google
    발신 메일 원문 자체의 표기 오류로 확인됨. 이 컴퓨터의 디코딩 문제가
    아님). 그래서 통화 기호는 특정 글자로 고정하지 않고, 숫자 앞의 글자
    하나를 그냥 건너뛰는 방식으로 처리한다."""
    # 상품명 안에 "(5 TB)"처럼 괄호 속 숫자가 섞여 있을 수 있어서, 금액은
    # 반드시 그 줄의 맨 끝에 오는 숫자여야 한다는 조건(줄바꿈 직전)을 걸어
    # 상품명 중간의 숫자를 금액으로 잘못 뽑는 걸 막는다. 연간 구독은 금액
    # 뒤에 "/연"처럼 단위가 붙기도 해서 그것도 허용한다.
    m = re.search(
        r"(?:상품\s*(?:가격|환불\s*금액)|Item\s+(?:Price|Amount\s*refunded))"
        r"\s*(.+?)\s+(?:매월\s*)?[^\d\s]?(\d[\d,]*)(?:/\S+)?[ \t]*(?:\r?\n|$)",
        body,
    )
    if not m:
        return "", None
    item_name = m.group(1).strip()
    amount = int(m.group(2).replace(",", ""))
    return item_name, amount


def extract_total_amount(body):
    # 통화 기호 표기 오류에 대응하는 이유는 extract_item_and_amount 참고.
    m = re.search(r"(?:합계|Total):\s*(?:매월\s*)?[^\d\s]?(\d[\d,]*)", body)
    return int(m.group(1).replace(",", "")) if m else None


def extract_payment_method(body):
    m = re.search(r"(?:결제\s*방법|Payment\s*method):\s*\n+\s*(\S[^\n]*)", body)
    return m.group(1).strip() if m else ""


def parse_order_datetime(text):
    """구글플레이 메일은 한글 형식(최근)과 영어 형식(예전) 둘 다 쓰인다."""
    # 한글: 2026. 9. 11. 오후 12시 50분 13초 GMT+9
    m = re.search(
        r"(\d{4})\.\s*(\d{1,2})\.\s*(\d{1,2})\.\s*(오전|오후)\s*"
        r"(\d{1,2})시\s*(\d{1,2})분\s*(\d{1,2})초",
        text,
    )
    if m:
        year, month, day = int(m.group(1)), int(m.group(2)), int(m.group(3))
        meridiem = m.group(4)
        hour, minute, second = int(m.group(5)), int(m.group(6)), int(m.group(7))
        if meridiem == "오후" and hour != 12:
            hour += 12
        if meridiem == "오전" and hour == 12:
            hour = 0
        try:
            return datetime(year, month, day, hour, minute, second)
        except ValueError:
            return None

    # 영어: Aug 9, 2025 1:25:27 AM PDT
    m = re.search(
        r"([A-Za-z]{3})\s+(\d{1,2}),\s*(\d{4})\s+(\d{1,2}):(\d{2}):(\d{2})\s*(AM|PM)",
        text,
    )
    if m and m.group(1) in MONTH_NAMES:
        month = MONTH_NAMES[m.group(1)]
        day, year = int(m.group(2)), int(m.group(3))
        hour, minute, second = int(m.group(4)), int(m.group(5)), int(m.group(6))
        meridiem = m.group(7)
        if meridiem == "PM" and hour != 12:
            hour += 12
        if meridiem == "AM" and hour == 12:
            hour = 0
        try:
            return datetime(year, month, day, hour, minute, second)
        except ValueError:
            return None

    return None


def parse_google_play_message(subject, body, msg_id, target_year, target_month, warn):
    # 제목에 "취소"(한글) 또는 "Cancellation"(예전 영어 형식)이 들어가면
    # 취소 영수증, 아니면 결제 영수증이다.
    is_cancel = "취소" in subject or "Cancellation" in subject

    date_text = extract_order_datetime_text(body)
    dt = parse_order_datetime(date_text)
    if dt is None:
        warn(f"메일(id={msg_id})에서 주문 날짜를 해석하지 못해 건너뜁니다.")
        return None
    if not (dt.year == target_year and dt.month == target_month):
        return None

    item_name, item_amount = extract_item_and_amount(body)
    amount = extract_total_amount(body)
    if amount is None:
        amount = item_amount
    if amount is None:
        warn(f"메일(id={msg_id})에서 결제금액을 찾지 못해 건너뜁니다.")
        return None

    return {
        "_dt": dt,
        "구분": "취소" if is_cancel else "결제",
        "결제일시": dt.strftime("%Y-%m-%d %H:%M"),
        "PG사": "Google Play",
        "상점명": extract_store_name(body),
        "상품명": item_name,
        "결제금액": amount,
        "주문번호": extract_order_number(body),
        "결제방법": extract_payment_method(body),
    }


def parse_google_payments_message(received_dt, body, target_year, target_month):
    """결제 센터(payments-noreply) "결제 완료" 메일. 본문 예:
        2026년 9월 15일에 ₩10,000의 결제 금액이 Google Cloud Platform & APIs에
        적용되었습니다.
    본문에 시각·주문번호·결제방법·상품명이 없어서, 결제일시는 메일 수신
    시각을 쓰고 나머지는 비워둔다. 월 판별도 수신 시각 기준으로 해서 엑셀의
    결제일시와 조회 월이 항상 일치하게 한다.
    이 문장이 없는 메일(결제 수단 업데이트, 계정 폐쇄 알림 등)은 돈이 움직인
    메일이 아니므로 경고 없이 제외한다."""
    if not (received_dt.year == target_year and received_dt.month == target_month):
        return None

    # 통화 기호 표기 오류에 대응하는 이유는 extract_item_and_amount 참고.
    m = re.search(
        r"[^\d\s]?(\d[\d,]*)의\s*결제\s*금액이\s*(.+?)에\s*적용",
        body,
        re.DOTALL,
    )
    if not m:
        return None

    return {
        "_dt": received_dt,
        "구분": "결제",
        "결제일시": received_dt.strftime("%Y-%m-%d %H:%M"),
        "PG사": "Google Payments",
        "상점명": re.sub(r"\s+", " ", m.group(2)).strip(),  # 줄바꿈된 서비스명을 한 줄로
        "상품명": "",
        "결제금액": int(m.group(1).replace(",", "")),
        "주문번호": "",
        "결제방법": "",
    }


# --- 5. 엑셀(xlsx) 저장 ------------------------------------------------------
def save_xlsx(rows, year, month):
    """정렬(왼쪽/오른쪽), 날짜·금액 서식, 헤더 필터가 적용된 엑셀 파일로 저장한다.
    주문번호는 숫자로 저장하면 지수 표기 등으로 깨질 수 있어 텍스트로 고정한다.
    결제일시는 내림차순으로 정렬해서 담는다."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"{year % 100:02d}{month:02d}00_구글플레이결제내역_claudecli.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "결제내역"

    ws.append(OUTPUT_COLUMNS)
    for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx).alignment = Alignment(
            horizontal="right" if col_name in RIGHT_ALIGN_COLUMNS else "left"
        )

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
            value = row["_dt"] if col_name == "결제일시" else row[col_name]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(
                horizontal="right" if col_name in RIGHT_ALIGN_COLUMNS else "left"
            )
            if col_name == "결제일시":
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif col_name == "결제금액":
                cell.number_format = '#,##0"원"'
            elif col_name == "주문번호":
                cell.number_format = "@"  # 텍스트 고정 (지수 표기 방지)

    last_row = len(rows) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{last_row}"

    wb.save(filepath)
    return filepath


# --- 메인 ------------------------------------------------------------------
def main():
    print("=== 구글플레이 월별 결제 내역 엑셀 만들기 ===\n")

    year, month = ask_year_month()
    print(f"\n{year}년 {month}월 결제 내역을 조회합니다. 잠시만 기다려주세요...\n")

    service = get_gmail_service()
    start, end = month_range(year, month)

    warnings = []

    def warn(message):
        warnings.append(message)
        print(f"  [건너뜀] {message}")

    rows = []

    query = build_query(start, end, GOOGLE_PLAY_SENDER, SUBJECT_FILTER)
    ids = list_message_ids(service, query)
    print(f"Google Play 대상 메일 {len(ids)}건")
    for msg_id in ids:
        subject, body, _ = get_message(service, msg_id)
        if not body:
            warn(f"메일(id={msg_id})의 내용을 읽을 수 없어 건너뜁니다.")
            continue
        row = parse_google_play_message(subject, body, msg_id, year, month, warn)
        if row:
            rows.append(row)

    query = build_query(start, end, GOOGLE_PAYMENTS_SENDER, PAYMENTS_SUBJECT_FILTER)
    ids = list_message_ids(service, query)
    print(f"Google 결제 센터 대상 메일 {len(ids)}건")
    for msg_id in ids:
        _, body, received_dt = get_message(service, msg_id)
        if not body:
            warn(f"메일(id={msg_id})의 내용을 읽을 수 없어 건너뜁니다.")
            continue
        row = parse_google_payments_message(received_dt, body, year, month)
        if row:
            rows.append(row)

    rows.sort(key=lambda r: r["_dt"], reverse=True)  # 결제일시 내림차순

    filepath = save_xlsx(rows, year, month)

    print()
    if rows:
        print(f"총 {len(rows)}건을 정리해서 저장했습니다.")
    else:
        print(f"{year}년 {month}월에는 구글플레이 결제 내역이 없습니다. (항목만 있는 빈 엑셀 파일을 만들었습니다)")
    print(f"저장 위치: {filepath}")

    if warnings:
        print(f"\n※ {len(warnings)}건은 형식을 해석하지 못해 건너뛰었습니다. 위 [건너뜀] 내용을 확인해주세요.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n취소되었습니다.")
    except PermissionError as e:
        print(f"\n[오류] 파일에 저장하지 못했습니다: {e.filename}")
        print("→ 이 엑셀 파일을 이미 열어놓은 상태라서 저장이 막혔을 가능성이 매우 높습니다.")
        print("   해당 파일을 닫은 뒤, 프로그램을 다시 실행해주세요.")
    except Exception as e:  # noqa: BLE001 - 비개발자용 프로그램이므로 원인을 그대로 보여준다
        print(f"\n[오류] 문제가 발생했습니다: {e}")
    finally:
        input("\n엔터 키를 누르면 창이 닫힙니다...")
